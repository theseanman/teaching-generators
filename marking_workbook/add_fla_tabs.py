#!/usr/bin/env python3
"""Add 'T1 Gradebook' and 'T1 Results' tabs for the FLA night class.

Header rows match the other gradebooks so populate_workbook.py needs no change:
  row 3 Category   (what 'Week' is for the day classes)
  row 4 Out of     (the mark total for a score column)
  row 5 Item       (the assignment title, as MARK.ONE names it)
  row 6 Outcome    ('level' or 'score')
  row 7 Applies to ('6 + 7', or '6'/'7' for one level only)
Students run rows 9-48, name in A, level 6 or 7 in B.

Results computes the percentage from the class's weighted categories the same way
MARK.ONE does — categories with nothing marked drop out and the rest scale up —
then the letter grade and proficiency from Sean's own band tables.
"""
import json, sys
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

FIRST, LAST = 9, 48
LEVELS = '"EM,DV,PR,EX,abs,ns"'
# Sean's FLA program conversion table
# ascending thresholds with MATCH type 1 = largest band at or below the percentage
GRADE = '{0,50,60,67,73,86}', '{"F","C-","C","C+","B","A"}'
PROF  = '{0,60,73,89}', '{"Emerging","Developing","Proficient","Extending"}'


def style_from(src_ws, coord):
    return copy(src_ws[coord]._style)


def build(wb, data, class_id='T1', label='FLA Term 1'):
    cls = next(c for c in data['classes'] if c['id'] == class_id)
    secs = [s for s in data['sections'] if s['classId'] == class_id]
    sec_level = {s['id']: s['level'] for s in secs}
    studs = []
    for s in data['students']:
        lv = next((sec_level[x] for x in s['sectionIds'] if x in sec_level), None)
        if lv:
            studs.append((lv, f"{s['last']}, {s['first']}"))
    studs.sort(key=lambda t: (t[0], t[1]))
    cats = [c['name'] for c in cls.get('categories', [])]
    weights = [c['weight'] for c in cls.get('categories', [])]
    asgs = [a for a in data['assignments']
            if a['classId'] == class_id and not a.get('recordOnly')]

    gsrc, rsrc = wb['A Gradebook'], wb['A Results']
    for name in (f'{class_id} Gradebook', f'{class_id} Results'):
        if name in wb.sheetnames:
            del wb[name]
    g = wb.create_sheet(f'{class_id} Gradebook')
    r = wb.create_sheet(f'{class_id} Results')

    # ---------------- Gradebook
    g['A1'] = f'{label} — gradebook (Literacy Foundations English, adult evening)'
    g['A1']._style = style_from(gsrc, 'A1')
    g['A2'] = ('Paste the roster into column A from row 9 and set each student\'s level '
               '(6 or 7) in B. Each assignment has a level column and a mark column. '
               'Levels: EM / DV / PR / EX, plus abs (absent) and ns (not submitted). '
               'The percentage, letter grade and proficiency are on the Results tab.')
    g['A2']._style = style_from(gsrc, 'A2')
    g['A3'] = 'Student (paste from the class list)'
    for row, lab in ((3, 'Category'), (4, 'Marked out of'), (5, 'Item'), (6, 'Outcome'), (7, 'Applies to')):
        g.cell(row, 2).value = lab
        g.cell(row, 2)._style = style_from(gsrc, f'B{row}')
        if row == 3:
            g.cell(row, 1)._style = style_from(gsrc, 'A3')

    col = 3
    asg_cols = []
    for a in asgs:
        for kind in ('level', 'score'):
            g.cell(3, col).value = a.get('category') or '—'
            g.cell(4, col).value = a.get('total') or ''
            g.cell(5, col).value = a.get('title', '')
            g.cell(6, col).value = kind
            g.cell(7, col).value = '6 + 7'
            for hr in (3, 4, 5, 6, 7):
                g.cell(hr, col)._style = style_from(gsrc, f'{L(4)}{hr}')
            g.column_dimensions[L(col)].width = 9 if kind == 'level' else 8
            asg_cols.append((a, kind, col))
            col += 1
    last_col = col - 1

    g.cell(8, 1).value = 'EXAMPLE — not counted'
    g.cell(8, 1)._style = style_from(gsrc, 'A8')
    g.cell(8, 2).value = '6'
    if asg_cols:
        g.cell(8, asg_cols[0][2]).value = 'PR'
        if len(asg_cols) > 1:
            g.cell(8, asg_cols[1][2]).value = 18
    for c in range(2, last_col + 1):
        g.cell(8, c)._style = style_from(gsrc, 'C8')

    for i, (lv, nm) in enumerate(studs):
        row = FIRST + i
        g.cell(row, 1).value = nm
        g.cell(row, 2).value = lv
        g.cell(row, 1)._style = style_from(gsrc, 'A9')
        g.cell(row, 2)._style = style_from(gsrc, 'B9')
        for c in range(3, last_col + 1):
            g.cell(row, c)._style = style_from(gsrc, 'D9')
    g.column_dimensions['A'].width = 26
    g.column_dimensions['B'].width = 8
    g.freeze_panes = 'C9'

    lv_cols = [L(c) for a, k, c in asg_cols if k == 'level']
    if lv_cols:
        dv = DataValidation(type='list', formula1=LEVELS, allow_blank=True)
        dv.sqref = ' '.join(f'{c}{FIRST}:{c}{LAST}' for c in lv_cols)
        g.add_data_validation(dv)
    dv6 = DataValidation(type='list', formula1='"6,7"', allow_blank=True)
    dv6.sqref = f'B{FIRST}:B{LAST}'
    g.add_data_validation(dv6)

    # ---------------- Results
    G = f"'{class_id} Gradebook'"
    r['A1'] = f'{label} — results'
    r['A1']._style = style_from(rsrc, 'A1')
    r['A2'] = ('Names, levels and marks come from the Gradebook. Percentage is the weighted '
               'average of the category averages, with categories that have nothing marked '
               'left out and the rest scaled up — the same rule MARK.ONE uses. Letter grade '
               'and proficiency follow the FLA program conversion table.')
    r['A2']._style = style_from(rsrc, 'A2')

    head = ['Student', 'Level'] + [f'{a.get("title","")} — level' for a, k, c in asg_cols if k == 'level'] \
           + [f'{a.get("title","")} — mark' for a, k, c in asg_cols if k == 'score'] \
           + ['Percentage', 'Letter grade', 'Proficiency', 'Interim proficiency']
    for i, h in enumerate(head, start=1):
        r.cell(4, i).value = h
        r.cell(4, i)._style = style_from(rsrc, 'C4' if i > 2 else f'{L(i)}4')
    n_lv = len(lv_cols)
    n_sc = len([1 for a, k, c in asg_cols if k == 'score'])
    c_pct = 3 + n_lv + n_sc
    c_let, c_prof, c_int = c_pct + 1, c_pct + 2, c_pct + 3
    h0 = c_int + 2                                    # helper band starts here

    for i in range(len(studs)):
        row = FIRST + i
        r.cell(row, 1).value = f'=IF({G}!A{row}="","",{G}!A{row})'
        r.cell(row, 2).value = f'=IF({G}!B{row}="","",{G}!B{row})'
        k = 3
        for a, kind, c in asg_cols:
            if kind == 'level':
                r.cell(row, k).value = f'=IF({G}!{L(c)}{row}="","",{G}!{L(c)}{row})'
                k += 1
        for a, kind, c in asg_cols:
            if kind == 'score':
                r.cell(row, k).value = (f'=IF({G}!{L(c)}{row}="","",'
                                        f'{G}!{L(c)}{row}&" / "&{G}!{L(c)}$4)')
                k += 1
        # helper band: per category, earned sum and the total actually attempted
        for j, cat in enumerate(cats):
            e, t = h0 + j * 2, h0 + j * 2 + 1
            r.cell(row, e).value = (f'=SUMIF({G}!$C$3:${L(last_col)}$3,"{cat}",'
                                    f'{G}!$C{row}:${L(last_col)}{row})')
            r.cell(row, t).value = (f'=SUMPRODUCT(({G}!$C$3:${L(last_col)}$3="{cat}")*'
                                    f'({G}!$C$6:${L(last_col)}$6="score")*'
                                    f'ISNUMBER({G}!$C{row}:${L(last_col)}{row})*'
                                    f'{G}!$C$4:${L(last_col)}$4)')
        num = '+'.join(f'IF({L(h0+j*2+1)}{row}>0,{weights[j]}*{L(h0+j*2)}{row}/'
                       f'{L(h0+j*2+1)}{row},0)' for j in range(len(cats)))
        den = '+'.join(f'IF({L(h0+j*2+1)}{row}>0,{weights[j]},0)' for j in range(len(cats)))
        r.cell(row, c_pct).value = (f'=IF($A{row}="","",IF(({den})=0,"",'
                                    f'ROUND(({num})/({den})*1000,0)/10))')
        pc = f'{L(c_pct)}{row}'
        r.cell(row, c_let).value = (f'=IF(OR($A{row}="",{pc}=""),"",'
                                    f'INDEX({GRADE[1]},MATCH({pc},{GRADE[0]},1)))')
        r.cell(row, c_prof).value = (f'=IF(OR($A{row}="",{pc}=""),"",'
                                     f'INDEX({PROF[1]},MATCH({pc},{PROF[0]},1)))')
        r.cell(row, c_int).value = f'=IF($A{row}="","",{L(c_prof)}{row})'
        for c in range(1, c_int + 1):
            r.cell(row, c)._style = style_from(rsrc, 'C9' if c > 2 else f'{L(c)}9')

    r.column_dimensions['A'].width = 26
    r.column_dimensions['B'].width = 7
    for c in range(3, c_int + 1):
        r.column_dimensions[L(c)].width = 19
    for j in range(len(cats) * 2):
        r.column_dimensions[L(h0 + j)].hidden = True
    r.freeze_panes = 'C9'
    return len(studs), len(asgs), cats


if __name__ == '__main__':
    src, datafile, out = sys.argv[1], sys.argv[2], sys.argv[3]
    wb = openpyxl.load_workbook(src)
    d = json.load(open(datafile))
    n, a, cats = build(wb, d)
    wb.save(out)
    print(f'T1 Gradebook + T1 Results built: {n} students, {a} assignments, '
          f'categories {cats}')
    print('tabs now:', wb.sheetnames)

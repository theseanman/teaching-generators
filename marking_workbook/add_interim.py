#!/usr/bin/env python3
"""Add an 'Interim proficiency' column to all six Results sheets.

ELL (A/E/C/G): converts Reported proficiency to the interim four-point wording,
   Not Yet Meeting -> Emerging, Minimally Meeting -> Developing,
   Meeting -> Developing, Fully Meeting -> Proficient. Extending is unused.
English 8 (F/H): already on the four-point scale, so it passes Reported proficiency through.

Writes into the single empty column between the visible block and the hidden helpers,
so nothing is inserted and no existing formula reference moves."""
import sys
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter as L

ELL_FROM = '{"Not Yet Meeting","Minimally Meeting","Meeting","Fully Meeting"}'
ELL_TO   = '{"Emerging","Developing","Developing","Proficient"}'
HEADER   = 'Interim proficiency'
FIRST, LAST = 9, 48


def layout(ws):
    """(last visible header column, the empty gap column, the reported-proficiency column)."""
    vis = [c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value is not None]
    last_vis = max(vis)
    rep = next(c for c in vis if str(ws.cell(4, c).value).strip() == 'Reported proficiency')
    gap = last_vis + 1
    used = any(ws.cell(r, gap).value not in (None, '') for r in range(1, LAST + 1))
    if used:
        raise SystemExit(f'{ws.title}: column {L(gap)} is not empty — stopping rather than overwrite')
    return last_vis, gap, rep


def main(src, out):
    wb = openpyxl.load_workbook(src)
    for blk in ['A', 'E', 'C', 'G', 'F', 'H']:
        name = f'{blk} Results'
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        last_vis, gap, rep = layout(ws)
        if str(ws.cell(4, gap).value or '') == HEADER:
            continue
        ell = blk in ('A', 'E', 'C', 'G')

        hdr_src, hdr_new = ws.cell(4, rep), ws.cell(4, gap)
        hdr_new.value = HEADER
        hdr_new._style = copy(hdr_src._style)

        for r in range(FIRST, LAST + 1):
            cell, model = ws.cell(r, gap), ws.cell(r, rep)
            rc = f'{L(rep)}{r}'
            cell.value = (
                f'=IF($A{r}="","",IFERROR(INDEX({ELL_TO},MATCH({rc},{ELL_FROM},0)),""))'
                if ell else f'=IF($A{r}="","",{rc})')
            cell._style = copy(model._style)

        ws.column_dimensions[L(gap)].width = max(
            18, ws.column_dimensions[L(rep)].width or 18)
        print(f'{name}: {HEADER} -> column {L(gap)} '
              f'({"converted from " + L(rep) if ell else "passes " + L(rep) + " through"})')
    wb.save(out)
    print('saved', out)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])

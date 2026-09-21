import sys
from pptx import Presentation
from pptx.util import Emu

def est_h(text, w_in, pt):
    if not text.strip(): return 0.0
    cpl = max(6, int(w_in / ((pt/72)*0.50)))
    lines = sum(max(1, -(-len(l)//cpl)) for l in text.split("\n"))
    return lines * (pt/72) * 1.22

p = Presentation(sys.argv[1])
SW, SH = p.slide_width/914400, p.slide_height/914400
bad = []
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if not t.strip(): continue
        try:
            w, h = sh.width/914400, sh.height/914400
            x, y = sh.left/914400, sh.top/914400
        except: continue
        pts = [r.font.size.pt for para in sh.text_frame.paragraphs for r in para.runs if r.font.size]
        pt = max(pts) if pts else 18
        need = est_h(t, w, pt)
        if need > h + 0.06:
            bad.append((i, "OVERFLOW", round(need,2), round(h,2), pt, t[:46].replace("\n"," ")))
        if y + h > SH + 0.05 or x + w > SW + 0.05:
            bad.append((i, "OFF-SLIDE", round(y+h,2), round(SH,2), pt, t[:46].replace("\n"," ")))
print(f"{len(p.slides)} slides, {len(bad)} problems")
for b in bad[:60]: print(f"  slide {b[0]:>3} {b[1]:<10} needs {b[2]}in in {b[3]}in @{b[4]}pt  \"{b[5]}\"")
